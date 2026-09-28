"""Assemble a document audit into one report: AUDIT_DIR/Audit Report - <label> - <id>.docx and .md.

Reads audit.json (scripted stage, `csa audit`) and sections/*.md (audit agent, `csa audit-section`).
Units without an agent file still appear, from the scripted checks alone.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

VERDICT_FILL = {"RED": "F8D7DA", "AMBER": "FFF3CD", "GREEN": "D4EDDA", "NOT AUDITED": "E2E3E5"}


def parse_section_file(p: Path) -> dict:
    t = p.read_text(encoding="utf-8")
    fm = dict(re.findall(r"^(\w+):\s*(.+)$", t.split("---")[1], re.M)) if t.startswith("---") else {}
    out = {"meta": fm, "summary": "", "tables": {}}
    m = re.search(r"## Summary\s*\n(.*?)(?=\n## |\Z)", t, re.S)
    out["summary"] = re.sub(r"\s+", " ", m.group(1)).strip() if m else ""
    for name in ("Coverage", "Claims", "Consistency", "Missing"):
        m = re.search(rf"## {name}\s*\n(.*?)(?=\n## |\Z)", t, re.S)
        rows = []
        if m:
            for line in m.group(1).splitlines():
                if line.startswith("|") and not re.match(r"^\|\s*-{3,}", line):
                    rows.append([c.strip() for c in line.strip().strip("|").split("|")])
        out["tables"][name] = rows
    return out


def scripted_flags(a: dict, title: str) -> list[str]:
    flags = []
    num = (re.match(r"^([\d.]+)\s", title) or [None, None])[1]
    for d in a["domains"]:
        if d["domain"] == title:
            flags += [f"missing {x}" for x in d.get("missing_subsections", [])]
    reqs = [r for r in a["requirements"]["rows"] if r["section"] == title or (num and r["section"].startswith(num + " "))]
    flags += [f"{r['req']}: {'; '.join(r['issues'])}" for r in reqs if r["issues"]]
    n_unk = sum(len(v) for k, v in a["unknowns"].items() if k == title or (num and k.startswith(num + ".")))
    if n_unk:
        flags.append(f"{n_unk} declared unknown(s)")
    n_ph = sum(1 for p in a["placeholders"] if p["section"] == title or (num and p["section"].startswith(num)))
    if n_ph:
        flags.append(f"{n_ph} placeholder(s)")
    n_c = sum(1 for c in a["document"].get("reviewer_comments", []) if c["section"] == title or (num and c["section"].startswith(num)))
    if n_c:
        flags.append(f"{n_c} open reviewer comment(s)")
    return flags


def build_markdown(a: dict, secs: dict) -> tuple[str, list[dict]]:
    units = [u for u in a["units"]]
    rows = []
    for u in units:
        s = secs.get(u["slug"])
        v = (s["meta"].get("verdict", "").split()[0].upper() if s and s["meta"].get("verdict") else "NOT AUDITED")
        rows.append({"unit": u["title"], "slug": u["slug"], "verdict": v, "coverage": s["meta"].get("coverage", "") if s else "",
                     "claims": s["meta"].get("claims", "") if s else "", "flags": scripted_flags(a, u["title"])})
    from collections import Counter
    vc = Counter(r["verdict"] for r in rows)
    d = a["document"]
    L = [f"# Document audit report", "", f"Project: {a.get('project') or '-'}  ", f"Document: {Path(a['docx']).name}  ",
         f"Audit: {a['audit_id']} ({a['generated_at']})", "", "## Summary", ""]
    L.append(f"Units: {len(rows)}; audited in depth: {len(rows) - vc.get('NOT AUDITED', 0)} "
             f"(RED {vc.get('RED', 0)}, AMBER {vc.get('AMBER', 0)}, GREEN {vc.get('GREEN', 0)}).")
    L.append(f"Requirements: {len(a['requirements']['rows'])} rows; missing from the document: {', '.join(a['requirements']['missing']) or 'none'}.")
    L.append(f"Open reviewer comments: {len(d.get('reviewer_comments', []))}. Tracked changes left: {d['tracked_insertions'] + d['tracked_deletions']}. "
             f"Placeholders: {len(a['placeholders'])}. Structure issues: {len(a['structure'])}. Appendix D items open: "
             f"{sum(1 for x in a['appendix_d'] if 'open' in x['status'].lower())} of {len(a['appendix_d'])}.")
    L += ["", "## Scorecard", "", "| Unit | Verdict | Coverage | Claims | Scripted findings |", "| --- | --- | --- | --- | --- |"]
    for r in rows:
        L.append(f"| {r['unit']} | {r['verdict']} | {r['coverage'] or '-'} | {r['claims'] or '-'} | {'; '.join(r['flags']) or '-'} |")
    L += ["", "## What is missing", "", "| Unit | Item | Why it matters | Owner / source | Action |", "| --- | --- | --- | --- | --- |"]
    n = 0
    for r in rows:
        s = secs.get(r["slug"])
        for m in (s["tables"].get("Missing", [])[1:] if s else []):
            if m and m[0].lower().rstrip(".") != "none":
                L.append(f"| {r['unit']} | " + " | ".join((m + ['', '', '', ''])[:4]) + " |"); n += 1
    for x in a["requirements"]["missing"]:
        L.append(f"| Requirements | {x} has no row in the document | Every checklist requirement needs a current state and rating | Author | Add the requirement row |"); n += 1
    for dm in a["domains"]:
        for mi in dm.get("missing_subsections", []):
            L.append(f"| {dm['domain']} | No {mi} | Every domain needs its discovery facts and isolation consequence | Author | Author the subsection |"); n += 1
    for s in a["structure"]:
        L.append(f"| Structure | {s['heading']} | {s['detail']} | Author | Fix the heading |"); n += 1
    for c in d.get("reviewer_comments", []):
        L.append(f"| {c['section']} | Reviewer comment: {c['text'][:120]} | Open review point | {c['author']} | Answer or resolve |"); n += 1
    for p in a["placeholders"]:
        L.append(f"| {p['section']} | Placeholder {p['text']} | Unfinished content | Author | Complete or remove |"); n += 1
    if n == 0:
        L.append("| - | None | - | - | - |")
    for r in rows:
        s = secs.get(r["slug"])
        L += ["", f"## {r['unit']}", "", f"Verdict: {r['verdict']}" + (f". Coverage {r['coverage']}, claims {r['claims']}." if s else ". Not audited in depth; scripted findings only."), ""]
        if s:
            if s["summary"]:
                L += [s["summary"], ""]
            for name in ("Coverage", "Claims", "Consistency", "Missing"):
                t = s["tables"].get(name) or []
                if len(t) > 1 and not (len(t) == 2 and t[1][0].lower().rstrip(".") == "none"):
                    L += [f"### {name}", "", "| " + " | ".join(t[0]) + " |", "| " + " | ".join("---" for _ in t[0]) + " |"]
                    L += ["| " + " | ".join(row) + " |" for row in t[1:]]
                    L.append("")
        if r["flags"]:
            L += ["Scripted findings: " + "; ".join(r["flags"]) + ".", ""]
    L += ["", "## Scripted checks", ""]
    mech = Path(a["out_dir"]) / "mechanical.md" if a.get("out_dir") else None
    return "\n".join(L) + "\n", rows


def write_docx(md: str, rows: list[dict], path: Path) -> str:
    try:
        import docx  # python-docx
        from docx.enum.section import WD_ORIENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError:
        return _write_docx_engine(md, path)
    d = docx.Document()
    sec = d.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.6))
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(9.5)

    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd")
        s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill); tcPr.append(s)

    lines = md.splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("|"):
            block = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|\s*-{3,}", lines[i]):
                    block.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            cols = max(len(r) for r in block)
            t = d.add_table(rows=0, cols=cols); t.style = "Table Grid"
            for ri, r in enumerate(block):
                cells = t.add_row().cells
                for ci in range(cols):
                    txt = r[ci] if ci < len(r) else ""
                    p = cells[ci].paragraphs[0]; run = p.add_run(txt); run.font.size = Pt(8.5)
                    if ri == 0:
                        run.bold = True; run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); shade(cells[ci], "1F3864")
                    elif txt in VERDICT_FILL:
                        shade(cells[ci], VERDICT_FILL[txt])
            d.add_paragraph()
            continue
        if ln.startswith("# "):
            d.add_heading(ln[2:], 0)
        elif ln.startswith("## "):
            d.add_heading(ln[3:], 1)
        elif ln.startswith("### "):
            d.add_heading(ln[4:], 2)
        elif ln.strip():
            d.add_paragraph(ln.rstrip(" "))
        i += 1
    d.save(str(path))
    return "python-docx"


def _write_docx_engine(md: str, path: Path) -> str:
    sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "vendor"))
    from docxengine._session import Session
    from docxengine._create import docx_create
    from docxengine._tools_lifecycle import docx_save
    s = Session(); r = docx_create(s, content_md=md)
    docx_save(s, doc_id=r["doc_id"], path=str(path))
    return "docxengine"


def build(audit_dir: Path) -> dict:
    a = json.loads((audit_dir / "audit.json").read_text(encoding="utf-8"))
    a["out_dir"] = str(audit_dir)
    secs = {}
    for f in sorted((audit_dir / "sections").glob("*.md")):
        secs[f.stem] = parse_section_file(f)
    md, rows = build_markdown(a, secs)
    md += (audit_dir / "mechanical.md").read_text(encoding="utf-8").split("\n", 2)[-1] if (audit_dir / "mechanical.md").is_file() else ""
    label = re.sub(r"[^\w -]", "", a.get("project") or "CSA")
    base = audit_dir / f"Audit Report - {label} - {a['audit_id']}"
    base.with_suffix(".md").write_text(md, encoding="utf-8")
    engine = write_docx(md, rows, base.with_suffix(".docx"))
    return {"status": "OK", "report": str(base.with_suffix(".docx")), "markdown": str(base.with_suffix(".md")), "engine": engine,
            "units": len(rows), "audited": sum(1 for r in rows if r["verdict"] != "NOT AUDITED")}


def main(argv=None) -> int:
    import argparse
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("audit_dir")
    a = ap.parse_args(argv)
    print(json.dumps(build(Path(a.audit_dir)), indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
