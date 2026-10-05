from __future__ import annotations

import csv
import hashlib
import json
import os
import re
import shutil
import sqlite3
import subprocess
import sys
import tempfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Literal

from fastapi import FastAPI, HTTPException, Query, Response
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from . import apply_workflow, audit_workflow, author_workflow, build_workflow, cleanup_workflow, commands, documents, pipeline, proposals, review_workflow, sections, validation, word_review
from .jobs import JobManager
from .models import CommandRequest, JobCreateRequest


APP_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_AGENTS_DIR = Path(__file__).resolve().parents[2]
SYSTEMS = {
    "iamps": {
        "label": "IAMPS",
        "application": "IAMPS",
        "project_key": "iamps-08",
        "scope": "The IAMPS application and the infrastructure, services and controls that operate it.",
    },
    "utcdtc": {
        "label": "UTC DTC",
        "application": "UTC DTC",
        "project_key": "utcdtc",
        "scope": "The UTC DTC application and the KVM-based system in which it operates.",
    },
    "tetra": {
        "label": "TETRA",
        "application": "Reveloc TETRA",
        "project_key": "tetra-reveloc",
        "scope": "The Reveloc TETRA application and the hosts, dependencies and controls that provide its service.",
    },
}
ANSWERED = "LIKELY_ANSWERED"
TEXT_EXTENSIONS = {
    "", ".txt", ".log", ".xml", ".csv", ".json", ".md", ".ps1", ".cfg",
    ".conf", ".ini", ".inf", ".bat", ".yaml", ".yml", ".html", ".htm",
}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg"}
OFFICE_EXTENSIONS = {".doc", ".docx", ".ppt", ".pptx"}
FAMILY_RULES = (
    ("Host", ("host", "systeminfo", "computer")),
    ("Time", ("time", "ntp", "w32tm")),
    ("Network", ("network", "route", "listener", "port", "firewall", "flow")),
    ("DNS", ("dns", "resolver")),
    ("Software", ("software", "application", "installed")),
    ("Services", ("service", "process")),
    ("Identity", ("account", "admin", "group", "user", "identity")),
    ("Policy", ("policy", "gpo", "gpresult")),
    ("Storage", ("disk", "volume", "storage", "share")),
)

AREA_LABELS = {
    "application_architecture": "Application architecture",
    "application_overview": "Application overview",
    "application_and_services": "Application and services",
    "architecture_and_dependencies": "Architecture and dependencies",
    "asset_inventory": "Assets and hosting",
    "availability_and_resilience": "Availability and resilience",
    "backup_and_recovery": "Backup and recovery",
    "business_and_operational_use": "Business and operational use",
    "business_continuity": "Continuity and resilience",
    "discovery_required": "Discovery required",
    "dns": "DNS",
    "gap_lookup": "Evidence candidates to review",
    "identity_and_access": "Identity and access",
    "identity_authentication": "Identity and authentication",
    "infrastructure_and_hosting": "Infrastructure and hosting",
    "integrations_and_dependencies": "Integrations and dependencies",
    "logging_and_monitoring": "Logging and monitoring",
    "monitoring_and_logging": "Monitoring and logging",
    "network_and_connectivity": "Network and connectivity",
    "operations_and_support": "Operations and support",
    "patching_and_vulnerability": "Patching and vulnerability",
    "pki_certificates": "PKI and certificates",
    "security_posture": "Security posture",
    "security_controls": "Security controls",
    "storage_and_data_transfer": "Storage and data transfer",
    "time_synchronisation": "Time synchronisation",
}

AREA_PRIORITY = (
    "application_overview", "business_and_operational_use", "application_architecture",
    "infrastructure_and_hosting", "application_and_services", "integrations_and_dependencies",
    "network_and_connectivity", "identity_and_access", "identity_authentication",
    "pki_certificates", "availability_and_resilience", "storage_and_data_transfer",
    "monitoring_and_logging", "logging_and_monitoring", "backup_and_recovery",
    "security_posture", "security_controls", "operations_and_support",
    "time_synchronisation", "dns", "discovery_required", "gap_lookup",
)

STATUS_ORDER = ("VERIFIED", "INFERRED", "UNCONFIRMED", "CONFLICTING", "NOT_FOUND")


class SearchRequest(BaseModel):
    query: str = Field(min_length=1, max_length=500)
    host: str | None = None
    include_archive: bool = False
    all_captures: bool = False


class EvidenceDraft(BaseModel):
    csa_area: str
    question: str
    claim: str
    status: Literal["VERIFIED", "INFERRED", "UNCONFIRMED", "CONFLICTING", "NOT_FOUND"]
    source_title: str = ""
    source_version: str = ""
    section: str = ""
    page_or_location: str = ""
    evidence_excerpt: str = ""
    inference_reason: str = ""
    confidence: Literal["", "high", "medium", "low"] = ""
    gap_or_action: str = ""
    review_state: str = "pending"


class CommitRequest(BaseModel):
    draft: EvidenceDraft
    confirmed: bool


def agents_dir() -> Path:
    return Path(os.environ.get("CSA_AGENTS_DIR", DEFAULT_AGENTS_DIR)).resolve()


def _parse_registry(path: Path) -> list[dict[str, str]]:
    if not path.is_file():
        raise HTTPException(503, f"Project registry is unavailable: {path}")
    items: list[dict[str, str]] = []
    current: dict[str, str] | None = None
    for raw in path.read_text(encoding="utf-8").splitlines():
        stripped = raw.strip()
        if not stripped or stripped.startswith("#") or stripped == "projects:":
            continue
        if stripped.startswith("- key:"):
            if current:
                items.append(current)
            current = {"key": stripped.split(":", 1)[1].strip().strip("\"'")}
        elif current is not None and ":" in stripped:
            key, value = stripped.split(":", 1)
            current[key.strip()] = value.strip().strip("\"'")
    if current:
        items.append(current)
    return items


def projects() -> dict[str, dict[str, str]]:
    path = Path(os.environ.get("CSA_PROJECTS_FILE", agents_dir() / "csa-context" / "PROJECTS.yaml"))
    return {p["key"]: p for p in _parse_registry(path)}


def system_project(system_key: str) -> dict[str, str]:
    if system_key not in SYSTEMS:
        raise HTTPException(404, "Unknown evidence system")
    project_key = SYSTEMS[system_key]["project_key"]
    project = projects().get(project_key)
    if not project:
        raise HTTPException(503, f"Registered project '{project_key}' is unavailable")
    return project


def work_dir(project: dict[str, str]) -> Path:
    return Path(project["work_dir"]).expanduser()


def open_index(project: dict[str, str]) -> sqlite3.Connection:
    db = work_dir(project) / "discovery-index.sqlite"
    if not db.is_file():
        raise HTTPException(503, f"Discovery index is unavailable: {db}")
    con = sqlite3.connect(f"file:{db}?immutable=1", uri=True)
    con.row_factory = sqlite3.Row
    return con


def run_json(command: list[str], *, stdin: str | None = None) -> dict[str, Any]:
    process = subprocess.run(command, input=stdin, capture_output=True, text=True)
    output = process.stdout.strip()
    try:
        data = json.loads(output) if output else {}
    except json.JSONDecodeError:
        raise HTTPException(502, process.stderr.strip() or output or "CSA helper returned invalid output")
    if process.returncode:
        detail = data.get("message") or data.get("error") or process.stderr.strip() or "CSA helper failed"
        raise HTTPException(422, {"message": detail, "result": data})
    return data


def matrix_script() -> Path:
    return agents_dir() / "skills" / "csa-evidence-matrix" / "scripts" / "evidence_matrix.py"


def index_script() -> Path:
    return agents_dir() / "skills" / "csa-discovery-index" / "scripts" / "discovery_index.py"


def reviews_for(project: dict[str, str]) -> dict[str, dict[str, Any]]:
    path = work_dir(project) / "evidence-reviews.jsonl"
    reviews: dict[str, dict[str, Any]] = {}
    if path.is_file():
        for line in path.read_text(encoding="utf-8").splitlines():
            if line.strip():
                row = json.loads(line)
                reviews[row["evidence_id"].upper()] = row
    return reviews


def overlay_review(row: dict[str, Any], reviews: dict[str, dict[str, Any]]) -> dict[str, Any]:
    review = reviews.get(str(row.get("evidence_id", "")).upper())
    if review:
        row = {**row, "review_state": review.get("state", "pending"),
               "reviewed_by": review.get("by"), "reviewed_at": review.get("at")}
    return row


def matrix_rows(project: dict[str, str]) -> list[dict[str, Any]]:
    matrix = work_dir(project) / "evidence-matrix.csv"
    if not matrix.is_file():
        return []
    reviews = reviews_for(project)
    with matrix.open(encoding="utf-8-sig", newline="") as handle:
        return [overlay_review(dict(row), reviews) for row in csv.DictReader(handle)]


def area_label(value: str) -> str:
    key = value.strip().lower()
    return AREA_LABELS.get(key, key.replace("_", " ").strip().title() or "Unclassified")


def scope_clause(all_captures: bool, include_archive: bool, alias: str = "f") -> tuple[str, list[Any]]:
    clauses: list[str] = []
    if not include_archive:
        clauses.append(f"{alias}.archived=0")
    if not all_captures:
        clauses.append(
            f"({alias}.capture_id IS NULL OR {alias}.capture_id IN "
            "(SELECT capture_id FROM captures WHERE superseded_by IS NULL))"
        )
    return " AND ".join(clauses) or "1=1", []


def family(value: str) -> str:
    lower = value.lower()
    for name, terms in FAMILY_RULES:
        if any(term in lower for term in terms):
            return name
    return "Other"


def source_roots(project: dict[str, str]) -> list[Path]:
    raw = [project.get("default_source_set", "")]
    raw.extend(project.get("index_extra_sources", "").split(";"))
    return [Path(p.strip()).expanduser().resolve() for p in raw if p.strip()]


def resolve_indexed_path(project: dict[str, str], rel_path: str) -> Path:
    """Resolve an indexed path, including indexes copied between CSA workspaces."""
    relative = Path(rel_path)
    direct = (Path(project["project_root"]) / relative).resolve()
    roots = source_roots(project)
    candidates = [direct]
    parts = relative.parts
    for root in roots:
        matches = [i for i, part in enumerate(parts) if part.casefold() == root.name.casefold()]
        if matches:
            candidates.append((root / Path(*parts[matches[-1] + 1:])).resolve())
    for candidate in candidates:
        if candidate.is_file() and any(candidate == root or root in candidate.parents for root in roots):
            return candidate
    capture = next((part for part in parts
                    if re.search(r"_discovery_[A-Za-z0-9-]+_\d{8}T\d{6}Z$", part, re.I)), None)
    recovered: list[Path] = []
    for root in roots:
        for match in root.rglob(relative.name):
            resolved = match.resolve()
            if capture and capture not in resolved.parts:
                continue
            recovered.append(resolved)
    unique = list(dict.fromkeys(recovered))
    if len(unique) == 1:
        return unique[0]
    return direct


def safe_file(project: dict[str, str], file_id: int) -> tuple[sqlite3.Row, Path]:
    con = open_index(project)
    row = con.execute("SELECT * FROM files WHERE file_id=?", (file_id,)).fetchone()
    con.close()
    if row is None:
        raise HTTPException(404, "Indexed file was not found")
    path = resolve_indexed_path(project, row["rel_path"])
    allowed = source_roots(project)
    if not any(path == root or root in path.parents for root in allowed):
        raise HTTPException(403, "Indexed path is outside configured evidence sources")
    if not path.is_file():
        raise HTTPException(404, f"Source file is unavailable: {path.name}")
    return row, path


def parse_anchor(value: str | None) -> tuple[int | None, int | None]:
    if not value:
        return None, None
    match = re.search(r"lines?\s+(\d+)(?:-(\d+))?", value, re.I)
    if not match:
        return None, None
    return int(match.group(1)), int(match.group(2) or match.group(1))


def text_preview(path: Path, anchor: str | None) -> dict[str, Any]:
    data = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    lines = (text or "").splitlines()
    start, end = parse_anchor(anchor)
    if start:
        first, last = max(1, start - 80), min(len(lines), (end or start) + 120)
    else:
        first, last = 1, min(len(lines), 2000)
    is_markup = path.suffix.lower() in {".xml", ".html", ".htm", ".xhtml"}
    return {
        "kind": "markup" if is_markup else "text", "format": path.suffix.lower().lstrip(".") if is_markup else None,
        "name": path.name, "line_start": first,
        "lines": [{"number": i, "text": lines[i - 1]} for i in range(first, last + 1)],
        "highlight": [start, end] if start else None,
        "truncated": last < len(lines), "total_lines": len(lines),
    }


def xlsx_preview(path: Path, anchor: str | None) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_name, target_row = workbook.sheetnames[0], 1
    if anchor:
        match = re.search(r"(?:sheet!row\s+)?([^!]+)!(\d+)", anchor, re.I)
        if match and match.group(1) in workbook.sheetnames:
            sheet_name, target_row = match.group(1), int(match.group(2))
    sheet = workbook[sheet_name]
    first = max(1, target_row - 30)
    last = min(sheet.max_row, target_row + 70)
    sheet_names = list(workbook.sheetnames)
    total_rows = sheet.max_row
    rows = [["" if cell.value is None else str(cell.value) for cell in row[:40]]
            for row in sheet.iter_rows(min_row=first, max_row=last)]
    workbook.close()
    return {"kind": "grid", "name": path.name, "sheets": sheet_names,
            "sheet": sheet_name, "row_start": first, "rows": rows,
            "highlight_row": target_row, "total_rows": total_rows}


def extracted_document_preview(project: dict[str, str], file_id: int, path: Path, anchor: str | None) -> dict[str, Any]:
    con = open_index(project)
    file_row = con.execute("SELECT coalesce(dup_of,file_id) source_id FROM files WHERE file_id=?", (file_id,)).fetchone()
    chunks = con.execute(
        "SELECT m.line_start, c.content FROM chunk_map m JOIN chunks c ON c.rowid=m.chunk_id "
        "WHERE m.file_id=? ORDER BY m.line_start", (file_row["source_id"],)
    ).fetchall()
    con.close()
    line_map: dict[int, str] = {}
    for chunk in chunks:
        for offset, value in enumerate(chunk["content"].splitlines()):
            tagged = re.match(r"\x1f(\d+)\x1f(.*)", value)
            number = int(tagged.group(1)) if tagged else chunk["line_start"] + offset
            line_map.setdefault(number, tagged.group(2) if tagged else value)
    start, end = parse_anchor(anchor)
    keys = sorted(line_map)
    if start:
        keys = [key for key in keys if max(1, start - 80) <= key <= (end or start) + 120]
    else:
        keys = keys[:2000]
    return {"kind": "document", "name": path.name,
            "lines": [{"number": key, "text": line_map[key]} for key in keys],
            "highlight": [start, end] if start else None}


def render_office(path: Path) -> Path:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise HTTPException(503, "LibreOffice preview runtime is unavailable")
    fingerprint = hashlib.sha256(f"{path}:{path.stat().st_mtime_ns}:{path.stat().st_size}".encode()).hexdigest()[:20]
    cache = Path(tempfile.gettempdir()) / "csa-evidence-preview" / fingerprint
    output = cache / f"{path.stem}.pdf"
    if output.is_file():
        return output
    cache.mkdir(parents=True, exist_ok=True)
    process = subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(cache), str(path)],
                             capture_output=True, text=True, timeout=120)
    if process.returncode or not output.is_file():
        raise HTTPException(502, process.stderr.strip() or "Document preview conversion failed")
    return output


app = FastAPI(title="CSA Evidence Workspace", version="0.1.0")
job_manager = JobManager(agents_dir())


@app.get("/api/systems")
def get_systems() -> dict[str, Any]:
    registered = projects()
    return {"systems": [{"key": key, "label": cfg["label"], "project_key": cfg["project_key"],
                         "available": cfg["project_key"] in registered} for key, cfg in SYSTEMS.items()]}


@app.get("/api/systems/{system_key}/commands")
def command_catalogue(system_key: str) -> dict[str, Any]:
    project = system_project(system_key)
    return {"system": system_key, "project_key": project["key"], "operations": commands.catalogue()}


@app.post("/api/systems/{system_key}/commands/{operation_key}")
def run_command(system_key: str, operation_key: str, request: CommandRequest) -> dict[str, Any]:
    project = system_project(system_key)
    return commands.execute(agents_dir(), project, operation_key, request).model_dump()


@app.post("/api/systems/{system_key}/jobs", status_code=202)
def create_job(system_key: str, request: JobCreateRequest) -> dict[str, Any]:
    project = system_project(system_key)
    command = CommandRequest(target=request.target, options=request.options)
    return job_manager.submit(project, request.operation, command).model_dump()


@app.get("/api/systems/{system_key}/jobs")
def list_jobs(system_key: str) -> dict[str, Any]:
    project = system_project(system_key)
    return {"jobs": [job.model_dump() for job in job_manager.list(project)]}


@app.get("/api/systems/{system_key}/jobs/{job_id}")
def get_job(system_key: str, job_id: str) -> dict[str, Any]:
    return job_manager.get(system_project(system_key), job_id).model_dump()


@app.delete("/api/systems/{system_key}/jobs/{job_id}")
def cancel_job(system_key: str, job_id: str) -> dict[str, Any]:
    return job_manager.cancel(system_project(system_key), job_id).model_dump()


@app.get("/api/systems/{system_key}/pipeline")
def get_pipeline(system_key: str, response: Response) -> dict[str, Any]:
    result = pipeline.snapshot(agents_dir(), system_project(system_key))
    response.headers["ETag"] = result.etag
    return result.model_dump()


@app.get("/api/systems/{system_key}/word-review")
def get_word_review(system_key: str) -> dict[str, Any]:
    return word_review.build_review(agents_dir(), system_project(system_key))


@app.post("/api/systems/{system_key}/word-review/open")
def open_word_review(system_key: str) -> dict[str, Any]:
    return word_review.open_authoritative(system_project(system_key))


@app.post("/api/systems/{system_key}/word-review/operator-note")
def record_operator_note(system_key: str, payload: dict) -> dict[str, Any]:
    note = str(payload.get("note", "")).strip()
    if not note:
        return {"state": "rejected", "reason": "note required"}
    return word_review.record_operator_note(agents_dir(), system_project(system_key), note)


@app.post("/api/systems/{system_key}/word-review/refresh")
def refresh_word_review(system_key: str) -> dict[str, Any]:
    return word_review.refresh_gate(agents_dir(), system_project(system_key))


@app.get("/api/systems/{system_key}/word-review/integrity")
def check_integrity(system_key: str, expected_hash: str | None = None) -> dict[str, Any]:
    return word_review.check_document_integrity(system_project(system_key), expected_hash)


@app.get("/api/systems/{system_key}/sections/{section}")
def get_section(system_key: str, section: str) -> dict[str, Any]:
    project = system_project(system_key)
    return sections.inspect(agents_dir(), project, section, matrix_rows(project))


@app.get("/api/systems/{system_key}/sections/{section}/proposal")
def get_section_proposal(system_key: str, section: str) -> dict[str, Any]:
    return proposals.load(agents_dir(), system_project(system_key), section)


@app.put("/api/systems/{system_key}/sections/{section}/proposal")
def save_section_proposal(system_key: str, section: str, payload: dict[str, Any]) -> dict[str, Any]:
    return proposals.save(agents_dir(), system_project(system_key), section, payload)


@app.post("/api/systems/{system_key}/sections/{section}/validate")
def validate_section_proposal(system_key: str, section: str, payload: dict[str, Any]) -> dict[str, Any]:
    return validation.validate(agents_dir(), system_project(system_key), section, payload)


@app.post("/api/systems/{system_key}/sections/{section}/apply-preflight")
def apply_preflight(system_key: str, section: str, payload: dict[str, Any]) -> dict[str, Any]:
    return apply_workflow.preflight(agents_dir(), system_project(system_key), section, payload)


@app.post("/api/systems/{system_key}/sections/{section}/apply", status_code=202)
def apply_section(system_key: str, section: str, payload: dict[str, Any]) -> dict[str, Any]:
    return apply_workflow.submit(agents_dir(), job_manager, system_project(system_key), section, payload)


@app.get("/api/systems/{system_key}/apply-jobs/{job_id}")
def get_apply_job(system_key: str, job_id: str) -> dict[str, Any]:
    job = job_manager.get(system_project(system_key), job_id).model_dump()
    return {"job": job, "outcome": apply_workflow.outcome(job)}


@app.post("/api/systems/{system_key}/sections/{section}/review", status_code=202)
def submit_review(system_key: str, section: str, payload: dict[str, Any]) -> dict[str, Any]:
    return review_workflow.submit(agents_dir(), job_manager, system_project(system_key), section, payload)


@app.get("/api/systems/{system_key}/review-jobs/{job_id}")
def get_review_job(system_key: str, job_id: str) -> dict[str, Any]:
    job = job_manager.get(system_project(system_key), job_id).model_dump()
    return {"job": job, "outcome": review_workflow.outcome(job)}


@app.get("/api/systems/{system_key}/review/latest/{section}")
def latest_review(system_key: str, section: str) -> dict[str, Any]:
    project = system_project(system_key)
    found = review_workflow.latest(agents_dir(), project, section)
    if found is None:
        raise HTTPException(404, "No completed review found for this section")
    return found


@app.get("/api/systems/{system_key}/sections/{section}/author-setup")
def author_setup(system_key: str, section: str) -> dict[str, Any]:
    project = system_project(system_key)
    return author_workflow.setup(agents_dir(), project, section, manager=job_manager)


@app.post("/api/systems/{system_key}/sections/{section}/author", status_code=202)
def submit_author(system_key: str, section: str, payload: dict[str, Any]) -> dict[str, Any]:
    unknown = set(payload) - {"cards"}
    if unknown:
        raise HTTPException(422, f"Unsupported author option(s): {', '.join(sorted(unknown))}")
    cards = payload.get("cards", False)
    if not isinstance(cards, bool):
        raise HTTPException(422, "cards must be true or false")
    return author_workflow.submit(agents_dir(), system_project(system_key), section,
                                  cards=cards, manager=job_manager)


@app.get("/api/systems/{system_key}/sections/{section}/author-jobs/{job_id}")
def get_author_job(system_key: str, section: str, job_id: str) -> dict[str, Any]:
    return author_workflow.get_job(agents_dir(), system_project(system_key), section, job_id,
                                   manager=job_manager)


@app.get("/api/systems/{system_key}/sections/{section}/author/latest")
def latest_author(system_key: str, section: str) -> dict[str, Any]:
    found = author_workflow.latest(agents_dir(), system_project(system_key), section, manager=job_manager)
    if found is None:
        raise HTTPException(404, "No author run found for this section")
    return found


@app.get("/api/systems/{system_key}/sections/{section}/author-artifacts/{artifact_key}")
def get_author_artifact(system_key: str, section: str, artifact_key: str) -> Response:
    content, media_type = author_workflow.artifact(system_project(system_key), section, artifact_key,
                                                   agents_dir=agents_dir())
    return Response(content=content, media_type=media_type,
                    headers={"Content-Disposition": f'inline; filename="{artifact_key}.txt"'})


@app.get("/api/systems/{system_key}/sections/{section}/cleanup-preflight")
def cleanup_preflight(system_key: str, section: str) -> dict[str, Any]:
    return cleanup_workflow.preflight(agents_dir(), system_project(system_key), section, {})


@app.post("/api/systems/{system_key}/sections/{section}/cleanup", status_code=202)
def submit_cleanup(system_key: str, section: str, payload: dict[str, Any]) -> dict[str, Any]:
    return cleanup_workflow.submit(agents_dir(), job_manager, system_project(system_key), section, payload)


@app.get("/api/systems/{system_key}/cleanup-jobs/{job_id}")
def get_cleanup_job(system_key: str, job_id: str) -> dict[str, Any]:
    job = job_manager.get(system_project(system_key), job_id).model_dump()
    return {"job": job, "outcome": cleanup_workflow.outcome(job)}


@app.get("/api/systems/{system_key}/cleanup/latest/{section}")
def latest_cleanup(system_key: str, section: str) -> dict[str, Any]:
    project = system_project(system_key)
    found = cleanup_workflow.latest(agents_dir(), project, section)
    if found is None:
        raise HTTPException(404, "No completed cleanup found for this section")
    return found


@app.get("/api/systems/{system_key}/build/setup")
def build_setup(system_key: str) -> dict[str, Any]:
    return build_workflow.setup(agents_dir(), system_project(system_key))


@app.post("/api/systems/{system_key}/build/preview", status_code=202)
def submit_build_preview(system_key: str) -> dict[str, Any]:
    return build_workflow.submit(agents_dir(), job_manager, system_project(system_key), preview=True, confirmed=True)


@app.post("/api/systems/{system_key}/build", status_code=202)
def submit_build(system_key: str, payload: dict[str, Any]) -> dict[str, Any]:
    return build_workflow.submit(agents_dir(), job_manager, system_project(system_key), preview=False,
                                 confirmed=payload.get("confirmed") is True)


@app.get("/api/systems/{system_key}/build-jobs/{job_id}")
def get_build_job(system_key: str, job_id: str) -> dict[str, Any]:
    job = job_manager.get(system_project(system_key), job_id).model_dump()
    return {"job": job, "outcome": build_workflow.outcome(job, preview=job.get("operation") == "build_preview")}


@app.get("/api/systems/{system_key}/audit")
def get_audit(system_key: str) -> dict[str, Any]:
    return audit_workflow.snapshot(system_project(system_key))


@app.post("/api/systems/{system_key}/audit", status_code=202)
def submit_audit(system_key: str) -> dict[str, Any]:
    return audit_workflow.submit(agents_dir(), job_manager, system_project(system_key))


@app.get("/api/systems/{system_key}/audit-jobs/{job_id}")
def get_audit_job(system_key: str, job_id: str) -> dict[str, Any]:
    project = system_project(system_key)
    job = job_manager.get(project, job_id).model_dump()
    return {"job": job, "outcome": audit_workflow.outcome(job, project)}


@app.post("/api/systems/{system_key}/audit/gaps")
def record_audit_gap(system_key: str, payload: dict[str, Any]) -> dict[str, Any]:
    return audit_workflow.record_gap(system_project(system_key), payload)


@app.get("/api/systems/{system_key}/document")
def get_document(system_key: str) -> dict[str, Any]:
    return documents.metadata(system_project(system_key))


@app.get("/api/systems/{system_key}/document/preview")
def get_document_preview(system_key: str, mode: str = "descriptor"):
    project = system_project(system_key)
    if mode == "revisions":
        return documents.revision_view(project)
    if mode == "pdf":
        path, fingerprint, _ = documents.render_pdf(project)
        return FileResponse(path, media_type="application/pdf", content_disposition_type="inline",
                            headers={"ETag": f'"{fingerprint}"'})
    if mode != "descriptor":
        raise HTTPException(400, "Preview mode must be descriptor, pdf or revisions")
    return documents.preview_descriptor(project, system_key)


@app.get("/api/systems/{system_key}/document/comments")
def get_document_comments(system_key: str) -> dict[str, Any]:
    return {"comments": documents.document_comments(system_project(system_key))}


@app.get("/api/systems/{system_key}/summary")
def get_summary(system_key: str) -> dict[str, Any]:
    project = system_project(system_key)
    con = open_index(project)
    meta = dict(con.execute("SELECT key,value FROM meta").fetchall())
    statuses = {row[0]: row[1] for row in con.execute("SELECT status,count(*) FROM files GROUP BY status")}
    captures = con.execute("SELECT count(*), sum(CASE WHEN superseded_by IS NULL AND archived=0 THEN 1 ELSE 0 END) FROM captures").fetchone()
    scope, _ = scope_clause(False, False)
    coverage: dict[str, dict[str, int]] = defaultdict(lambda: defaultdict(int))
    capture_rows = con.execute(
        f"SELECT coalesce(f.host,'Documents') host, f.rel_path FROM files f WHERE {scope} AND f.status='indexed'"
    ).fetchall()
    for row in capture_rows:
        coverage[row["host"]][family(row["rel_path"])] += 1
    matrix = work_dir(project) / "evidence-matrix.csv"
    evidence_count = 0
    if matrix.is_file():
        with matrix.open(encoding="utf-8-sig", newline="") as handle:
            evidence_count = sum(1 for _ in csv.DictReader(handle))
    con.close()
    families = [name for name, _ in FAMILY_RULES] + ["Other"]
    return {"system": system_key, "label": SYSTEMS[system_key]["label"], "project_key": project["key"],
            "built_at_utc": meta.get("built_at_utc"), "index_complete": meta.get("complete") == "1",
            "captures": captures[0] or 0, "current_captures": captures[1] or 0,
            "files_by_status": statuses, "evidence_rows": evidence_count, "families": families,
            "coverage": [{"host": host, "families": dict(values)} for host, values in sorted(coverage.items())]}


@app.get("/api/systems/{system_key}/portrait")
def application_portrait(system_key: str) -> dict[str, Any]:
    """Translate indexed material into the questions an assessor must resolve."""
    project = system_project(system_key)
    config = SYSTEMS[system_key]
    rows = matrix_rows(project)
    con = open_index(project)
    scope, params = scope_clause(False, False)
    files = con.execute(
        "SELECT f.status,f.reason,f.host,f.rel_path FROM files f WHERE " + scope, params
    ).fetchall()
    con.close()

    status_counts = Counter((row.get("status") or "UNCONFIRMED").upper() for row in rows)
    review_counts = Counter((row.get("review_state") or "pending").lower() for row in rows)
    areas: list[dict[str, Any]] = []
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        grouped[row.get("csa_area", "")].append(row)

    for key, area_rows in grouped.items():
        counts = Counter((row.get("status") or "UNCONFIRMED").upper() for row in area_rows)
        open_rows = [row for row in area_rows if (row.get("status") or "").upper() != "VERIFIED"]
        verified = [row for row in area_rows if (row.get("status") or "").upper() == "VERIFIED"]
        if counts["CONFLICTING"]:
            state = "conflict"
        elif counts["NOT_FOUND"] or counts["UNCONFIRMED"]:
            state = "gaps"
        elif counts["INFERRED"]:
            state = "inference"
        elif verified:
            state = "supported"
        else:
            state = "empty"
        areas.append({
            "key": key or "unclassified",
            "label": area_label(key),
            "state": state,
            "total": len(area_rows),
            "counts": {status: counts.get(status, 0) for status in STATUS_ORDER},
            "pending_review": sum(1 for row in area_rows if (row.get("review_state") or "pending").lower() == "pending"),
            "claims": [{field: row.get(field, "") for field in
                        ("evidence_id", "claim", "status", "confidence", "source_title", "page_or_location")}
                       for row in verified[:4]],
            "decisions": [{field: row.get(field, "") for field in
                           ("evidence_id", "question", "claim", "status", "confidence", "gap_or_action")}
                          for row in open_rows[:8]],
        })
    priority = {key: index for index, key in enumerate(AREA_PRIORITY)}
    areas.sort(key=lambda area: (priority.get(area["key"], len(priority)), area["label"]))

    host_families: dict[str, Counter[str]] = defaultdict(Counter)
    source_status = Counter()
    skipped_reasons = Counter()
    for row in files:
        source_status[row["status"]] += 1
        if row["status"] == "skipped":
            skipped_reasons[row["reason"] or "No reason recorded"] += 1
        if row["host"] and row["status"] == "indexed":
            host_families[row["host"]][family(row["rel_path"])] += 1
    observed_hosts = [{
        "host": host,
        "families": dict(counts),
        "files": sum(counts.values()),
        "breadth": len(counts),
    } for host, counts in sorted(host_families.items())]

    attention = [row for row in rows if (row.get("status") or "").upper() in
                 {"CONFLICTING", "NOT_FOUND", "UNCONFIRMED", "INFERRED"}]
    attention.sort(key=lambda row: (
        {"CONFLICTING": 0, "NOT_FOUND": 1, "UNCONFIRMED": 2, "INFERRED": 3}.get(
            (row.get("status") or "").upper(), 4),
        row.get("evidence_id", ""),
    ))
    return {
        "application": {
            "name": config["application"],
            "system": config["label"],
            "project_key": project["key"],
            "scope": config["scope"],
        },
        "assessment": {
            "total": len(rows),
            "status_counts": {status: status_counts.get(status, 0) for status in STATUS_ORDER},
            "review_counts": dict(review_counts),
            "supported_areas": sum(1 for area in areas if area["state"] == "supported"),
            "areas_with_attention": sum(1 for area in areas if area["state"] in {"conflict", "gaps", "inference"}),
        },
        "areas": areas,
        "attention": [{field: row.get(field, "") for field in
                       ("evidence_id", "csa_area", "question", "claim", "status", "confidence",
                        "gap_or_action", "source_title", "page_or_location")}
                      for row in attention],
        "observed_hosts": observed_hosts,
        "source_health": {
            "files": sum(source_status.values()),
            "status_counts": dict(source_status),
            "skipped_reasons": [{"reason": reason, "count": count}
                                for reason, count in skipped_reasons.most_common(8)],
        },
        "disclaimer": "This portrait summarises recorded evidence. Host coverage shows where data was captured; it does not by itself prove an application dependency.",
    }


@app.post("/api/systems/{system_key}/search")
def search(system_key: str, request: SearchRequest) -> dict[str, Any]:
    project = system_project(system_key)
    matrix = run_json([sys.executable, str(matrix_script()), "--workspace", project["project_root"],
                       "lookup", request.query, "--limit", "12"])
    reviews = reviews_for(project)
    matrix["matches"] = [overlay_review(row, reviews) for row in matrix.get("matches", [])]
    stages: list[dict[str, Any]] = [{"kind": "matrix", "state": "complete", **matrix}]
    if matrix.get("verdict") != ANSWERED:
        command = [sys.executable, str(index_script()), "--project", project["key"], "search", request.query,
                   "--limit", "40", "--per-file", "3"]
        if request.host:
            command += ["--host", request.host]
        if request.include_archive:
            command.append("--include-archive")
        if request.all_captures:
            command.append("--all-captures")
        index = run_json(command)
        con = open_index(project)
        for hit in index.get("results", []):
            file_row = con.execute("SELECT file_id FROM files WHERE rel_path=?", (hit.get("rel_path"),)).fetchone()
            hit["file_id"] = file_row[0] if file_row else None
        con.close()
        stages.append({"kind": "index", "state": "complete", **index})
    return {"system": system_key, "query": request.query, "stages": stages}


@app.get("/api/systems/{system_key}/evidence")
def evidence(system_key: str, q: str = "", status: str = "", area: str = "", offset: int = 0,
             limit: int = Query(100, ge=1, le=500)) -> dict[str, Any]:
    project = system_project(system_key)
    matrix = work_dir(project) / "evidence-matrix.csv"
    if not matrix.is_file():
        raise HTTPException(404, "Evidence matrix is unavailable")
    reviews = reviews_for(project)
    with matrix.open(encoding="utf-8-sig", newline="") as handle:
        rows = [overlay_review(dict(row), reviews) for row in csv.DictReader(handle)]
    if q:
        words = q.casefold().split()
        rows = [row for row in rows if all(word in " ".join(row.values()).casefold() for word in words)]
    if status:
        rows = [row for row in rows if row.get("status") == status]
    if area:
        rows = [row for row in rows if row.get("csa_area") == area]
    return {"total": len(rows), "rows": rows[offset:offset + limit],
            "areas": sorted({row.get("csa_area", "") for row in rows if row.get("csa_area")})}


@app.get("/api/systems/{system_key}/files")
def files(system_key: str, q: str = "", host: str = "", ext: str = "", status: str = "", source_class: str = "",
          all_captures: bool = False, include_archive: bool = False, offset: int = 0,
          limit: int = Query(150, ge=1, le=500)) -> dict[str, Any]:
    project = system_project(system_key)
    con = open_index(project)
    scope, params = scope_clause(all_captures, include_archive)
    clauses = [scope]
    if q:
        clauses.append("(f.rel_path LIKE ? OR coalesce(f.host,'') LIKE ?)")
        params.extend([f"%{q}%", f"%{q}%"])
    for field, value in (("host", host), ("ext", ext), ("status", status), ("source_class", source_class)):
        if value:
            clauses.append(f"f.{field} LIKE ?")
            params.append(f"%{value}%")
    where = " AND ".join(clauses)
    total = con.execute(f"SELECT count(*) FROM files f WHERE {where}", params).fetchone()[0]
    rows = [dict(row) for row in con.execute(
        "SELECT f.file_id,f.rel_path,f.host,f.ext,f.size,f.status,f.reason,f.source_class,f.archived,"
        "c.capture_utc,c.superseded_by,CASE WHEN c.capture_id IS NULL OR c.superseded_by IS NULL THEN 1 ELSE 0 END current "
        f"FROM files f LEFT JOIN captures c ON c.capture_id=f.capture_id WHERE {where} "
        "ORDER BY coalesce(f.host,''),f.rel_path LIMIT ? OFFSET ?", params + [limit, offset]
    )]
    facets = {
        "statuses": [row[0] for row in con.execute("SELECT DISTINCT status FROM files ORDER BY status")],
        "extensions": [row[0] for row in con.execute("SELECT DISTINCT ext FROM files ORDER BY ext")],
        "source_classes": [row[0] for row in con.execute("SELECT DISTINCT source_class FROM files ORDER BY source_class")],
    }
    con.close()
    return {"total": total, "files": rows, "facets": facets}


@app.get("/api/systems/{system_key}/files/{file_id}/preview")
def preview(system_key: str, file_id: int, anchor: str | None = None, mode: str = "preview"):
    project = system_project(system_key)
    row, path = safe_file(project, file_id)
    ext = path.suffix.lower()
    if mode == "raw":
        return FileResponse(path, filename=path.name, content_disposition_type="inline")
    if mode == "pdf":
        if ext == ".pdf":
            return FileResponse(path, media_type="application/pdf", content_disposition_type="inline")
        if ext in OFFICE_EXTENSIONS:
            return FileResponse(render_office(path), media_type="application/pdf", content_disposition_type="inline")
        raise HTTPException(400, "This file does not have a PDF preview")
    if ext in TEXT_EXTENSIONS:
        result = text_preview(path, anchor)
    elif ext == ".xlsx":
        result = xlsx_preview(path, anchor)
    elif ext == ".pdf" or ext in OFFICE_EXTENSIONS:
        result = extracted_document_preview(project, file_id, path, anchor)
        result["pdf_url"] = f"/api/systems/{system_key}/files/{file_id}/preview?mode=pdf"
    elif ext in IMAGE_EXTENSIONS:
        result = {"kind": "image", "name": path.name,
                  "raw_url": f"/api/systems/{system_key}/files/{file_id}/preview?mode=raw"}
    else:
        result = {"kind": "unsupported", "name": path.name, "reason": row["reason"],
                  "raw_url": f"/api/systems/{system_key}/files/{file_id}/preview?mode=raw"}
    result.update({"file_id": file_id, "rel_path": row["rel_path"], "status": row["status"],
                   "reason": row["reason"], "size": row["size"]})
    return result


@app.get("/api/systems/{system_key}/evidence/{evidence_id}/sources")
def evidence_sources(system_key: str, evidence_id: str) -> dict[str, Any]:
    project = system_project(system_key)
    data = evidence(system_key, q="", limit=500)
    target = next((row for row in data["rows"] if row["evidence_id"].upper() == evidence_id.upper()), None)
    if not target:
        raise HTTPException(404, "Evidence row was not found")
    names = re.findall(r"[\w.\-]+\.(?:txt|csv|xlsx|docx|doc|pdf|md|log|json|xml|ps1|pptx)\b",
                       target.get("source_title", ""), re.I)
    con = open_index(project)
    matches: list[dict[str, Any]] = []
    current_scope, scope_params = scope_clause(False, False)
    for name in sorted(set(names)):
        rows = con.execute("SELECT f.file_id,f.rel_path,f.host,f.status FROM files f "
                           f"WHERE f.rel_path LIKE ? AND {current_scope}",
                           [f"%/{name}", *scope_params]).fetchall()
        matches.extend(dict(row) for row in rows)
    con.close()
    citation = " ".join((target.get("source_title", ""), target.get("source_version", "")))
    capture_ids = set(re.findall(r"[A-Za-z0-9-]+_discovery_[A-Za-z0-9-]+_\d{8}T\d{6}Z", citation, re.I))
    capture_ids.update(re.findall(r"[A-Za-z0-9-]+_\d{8}T\d{6}Z", citation, re.I))
    if capture_ids:
        scoped = [row for row in matches if any(capture.casefold() in row["rel_path"].casefold()
                                                for capture in capture_ids)]
        if scoped:
            matches = scoped
    else:
        cited_hosts = {str(row.get("host", "")).casefold() for row in matches
                       if row.get("host") and re.search(rf"\b{re.escape(str(row['host']))}\b", citation, re.I)}
        if cited_hosts:
            matches = [row for row in matches if str(row.get("host", "")).casefold() in cited_hosts]
    matches = list({row["file_id"]: row for row in matches}.values())
    return {"evidence": target, "sources": matches,
            "state": "found" if len(matches) == 1 else "ambiguous" if matches else "unavailable"}


@app.get("/api/systems/{system_key}/tables")
def tables(system_key: str, table: str = "", host: str = "", where: str = "", limit: int = 100) -> dict[str, Any]:
    project = system_project(system_key)
    if not table:
        return run_json([sys.executable, str(index_script()), "--project", project["key"], "tables"])
    command = [sys.executable, str(index_script()), "--project", project["key"], "rows", table,
               "--limit", str(min(max(limit, 1), 500))]
    if host:
        command += ["--host", host]
    if where:
        command += ["--where", where]
    return run_json(command)


def evidence_command(project: dict[str, str], draft: EvidenceDraft, dry_run: bool) -> dict[str, Any]:
    command = [sys.executable, str(matrix_script()), "--workspace", project["project_root"], "append",
               "--agent", "csa evidence workspace", "--context", "browser evidence review", "--rows-file", "-"]
    if dry_run:
        command.append("--dry-run")
    return run_json(command, stdin=json.dumps(draft.model_dump(), ensure_ascii=False))


@app.post("/api/systems/{system_key}/evidence-drafts/validate")
def validate_evidence(system_key: str, draft: EvidenceDraft) -> dict[str, Any]:
    return evidence_command(system_project(system_key), draft, True)


@app.post("/api/systems/{system_key}/evidence-drafts/commit")
def commit_evidence(system_key: str, request: CommitRequest) -> dict[str, Any]:
    if not request.confirmed:
        raise HTTPException(409, "Evidence append requires explicit confirmation")
    return evidence_command(system_project(system_key), request.draft, False)


DIST = APP_ROOT / "frontend" / "dist"
if DIST.is_dir():
    app.mount("/", StaticFiles(directory=DIST, html=True), name="workspace")
